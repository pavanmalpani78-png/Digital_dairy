from pptx import Presentation
from pptx.util import Inches, Pt

# Initialize presentation
prs = Presentation()

# --- Slide 1: Title Slide ---
slide_layout_title = prs.slide_layouts[0] 
slide_1 = prs.slides.add_slide(slide_layout_title)
title_1 = slide_1.shapes.title
subtitle_1 = slide_1.placeholders[1]

title_1.text = "Digital Dairy"
subtitle_1.text = "Web-Based Milk Order & Billing Management System\n\nPresented By: [Student Names]\nGuide: [Guide Name]\nCollege: [College Name]"

# --- Slide 2: Introduction ---
slide_layout_bullet = prs.slide_layouts[1] 
slide_2 = prs.slides.add_slide(slide_layout_bullet)
title_2 = slide_2.shapes.title
content_2 = slide_2.placeholders[1]

title_2.text = "Introduction"
tf_2 = content_2.text_frame
tf_2.text = "Dairy industry plays an important role in India."
p = tf_2.add_paragraph()
p.text = "Small dairy shops handle daily milk supply for millions of households."
p = tf_2.add_paragraph()
p.text = "Digital transformation is needed to modernize these traditional businesses."

# --- Slide 3: Background ---
slide_3 = prs.slides.add_slide(slide_layout_bullet)
title_3 = slide_3.shapes.title
content_3 = slide_3.placeholders[1]

title_3.text = "Background"
tf_3 = content_3.text_frame
tf_3.text = "Most dairy shops currently rely heavily on manual paper registers."
p = tf_3.add_paragraph()
p.text = "Daily entries consist of milk quantities, morning/evening shifts, and partial payments."
p = tf_3.add_paragraph()
p.text = "As the customer base grows, this data becomes increasingly difficult to manage."

# --- Slide 4: Problem Statement ---
slide_4 = prs.slides.add_slide(slide_layout_bullet)
title_4 = slide_4.shapes.title
content_4 = slide_4.placeholders[1]

title_4.text = "Problem Statement"
tf_4 = content_4.text_frame
tf_4.text = "Manual system is highly time-consuming."
p = tf_4.add_paragraph()
p.text = "High chances of human errors in calculations and billing."
p = tf_4.add_paragraph()
p.text = "No proper data security or reliable data backup."

# --- Slide 5: Need of Project ---
slide_5 = prs.slides.add_slide(slide_layout_bullet)
title_5 = slide_5.shapes.title
content_5 = slide_5.placeholders[1]

title_5.text = "Need of Project"
tf_5 = content_5.text_frame
tf_5.text = "Reduce tedious manual daily work."
p = tf_5.add_paragraph()
p.text = "Improve accuracy in complex monthly billing calculations."
p = tf_5.add_paragraph()
p.text = "Enable fast data retrieval for customer history."
p = tf_5.add_paragraph()
p.text = "Ensure secure digital storage to prevent record loss."

# Save the presentation
prs.save('Digital_Dairy_Presentation.pptx')
print("PowerPoint presentation 'Digital_Dairy_Presentation.pptx' created successfully!")